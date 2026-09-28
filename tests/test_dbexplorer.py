"""Run with:  python3 -m unittest discover tests

The Excel round-trip checks use openpyxl if it's installed (pip install
openpyxl); everything else is standard library.
"""

import io
import json
import re
import os
import shutil
import sqlite3
import sys
import tempfile
import threading
import time
import unittest
import urllib.error
import urllib.request
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import dbexplorer  # noqa: E402
import xlsx_writer  # noqa: E402
from xlsx_writer import XlsxWriter  # noqa: E402

try:
    import openpyxl
except ImportError:
    openpyxl = None

SAMPLE = os.path.join(ROOT, "examples", "sample.sqlite")


def write_xlsx(sheets):
    buf = io.BytesIO()
    w = XlsxWriter(buf)
    for name, cols, rows in sheets:
        w.add_sheet(name, cols, rows)
    w.close()
    buf.seek(0)
    return buf, w


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class ExcelValuesSurvive(unittest.TestCase):
    def test_values_come_back_exactly(self):
        rows = [
            ("00123", 42, 9007199254740993, 123.456, None, "2024-01-01", "=SUM(A1:A2)", b"\x01\xff",
             "_x0041_ stays", "tab\there\x01bell", 12345678901234, -5.5e-10),
        ]
        cols = ["zip", "n", "big", "real", "null", "date", "formula", "blob", "escape", "ctrl", "long", "tiny"]
        buf, w = write_xlsx([("t", cols, rows)])
        ws = openpyxl.load_workbook(buf).active
        got = [c.value for c in ws[2]]
        self.assertEqual(got[0], "00123")
        self.assertEqual(got[1], 42)
        self.assertEqual(got[2], "9007199254740993")  # >15 digits -> text, every digit kept
        self.assertEqual(got[3], 123.456)
        self.assertIsNone(got[4])
        self.assertEqual(got[5], "2024-01-01")        # not converted to a date
        self.assertEqual(got[6], "=SUM(A1:A2)")       # not a formula
        self.assertEqual(ws.cell(2, 7).data_type, "s")
        self.assertEqual(got[7], "0x01ff")
        # Excel decodes _xHHHH_ escapes (openpyxl doesn't), so check the raw XML:
        # the literal "_x0041_" is protected and the control character escaped.
        with zipfile.ZipFile(buf) as z:
            xml = z.read("xl/worksheets/sheet1.xml").decode()
        self.assertIn("_x005F_x0041_ stays", xml)
        self.assertIn("tab\there_x0001_bell", xml)
        self.assertEqual(got[10], 12345678901234)
        self.assertEqual(ws.cell(2, 11).number_format, "0")  # shown in full, not 1.23E+13
        self.assertEqual(got[11], -5.5e-10)
        self.assertEqual(ws.cell(2, 1).number_format, "@")   # text stays text when edited
        self.assertEqual(ws.freeze_panes, "A2")
        self.assertEqual(ws.auto_filter.ref, "A1:L2")
        self.assertEqual(w.warnings["big_ints"], 1)

    def test_long_results_continue_on_extra_sheets(self):
        old = xlsx_writer.MAX_ROWS
        xlsx_writer.MAX_ROWS = 4  # header + 3 rows per sheet
        try:
            buf, w = write_xlsx([("data", ["n"], ((i,) for i in range(7)))])
        finally:
            xlsx_writer.MAX_ROWS = old
        wb = openpyxl.load_workbook(buf)
        self.assertEqual(wb.sheetnames, ["data", "data (2)", "data (3)"])
        values = [c.value for ws in wb for c in ws["A"][1:]]
        self.assertEqual(values, list(range(7)))
        self.assertEqual(w.rows_written, 7)

    def test_exact_multiple_of_limit_makes_no_empty_sheet(self):
        old = xlsx_writer.MAX_ROWS
        xlsx_writer.MAX_ROWS = 4
        try:
            buf, _ = write_xlsx([("data", ["n"], ((i,) for i in range(6)))])
        finally:
            xlsx_writer.MAX_ROWS = old
        self.assertEqual(openpyxl.load_workbook(buf).sheetnames, ["data", "data (2)"])

    def test_sheet_names_are_made_valid_and_unique(self):
        buf, _ = write_xlsx([("a/b:c*?[x]", ["x"], []), ("A/B:C*?[X]", ["x"], []), ("x" * 40, ["x"], [])])
        names = openpyxl.load_workbook(buf).sheetnames
        self.assertEqual(names[0], "a_b_c___x_")
        self.assertEqual(names[1], "A_B_C___X_ (2)")
        self.assertEqual(len(names[2]), 31)


class ExcelStructure(unittest.TestCase):
    def test_is_a_valid_package(self):
        buf, _ = write_xlsx([("t", ["a"], [("x",)])])
        with zipfile.ZipFile(buf) as z:
            self.assertIsNone(z.testzip())
            self.assertIn("xl/worksheets/sheet1.xml", z.namelist())
            self.assertIn("xl/workbook.xml", z.namelist())

    def test_empty_workbook_still_has_a_sheet(self):
        buf = io.BytesIO()
        w = XlsxWriter(buf)
        w.close()
        with zipfile.ZipFile(buf) as z:
            self.assertIn("xl/worksheets/sheet1.xml", z.namelist())


class ReadOnly(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.path = os.path.join(self.dir, "t.db")
        shutil.copy(SAMPLE, self.path)

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def test_writes_are_refused(self):
        conn = dbexplorer.connect(self.path)
        for sql in ["DELETE FROM regions", "CREATE TABLE x(a)", "CREATE TEMP TABLE x(a)",
                    "PRAGMA query_only = 0", f"ATTACH '{self.dir}/other.db' AS o"]:
            with self.assertRaises(sqlite3.DatabaseError, msg=sql):
                conn.execute(sql)
        self.assertFalse(os.path.exists(os.path.join(self.dir, "other.db")))
        self.assertEqual(conn.execute("SELECT COUNT(*) FROM regions").fetchone()[0], 6)

    def test_schema_reading_is_allowed(self):
        tables = dbexplorer.load_schema(dbexplorer.connect(self.path))
        names = {t["name"] for t in tables}
        self.assertTrue({"customers", "orders", "order_items", "order_totals"} <= names)
        items = next(t for t in tables if t["name"] == "order_items")
        self.assertEqual({fk["table"] for fk in items["fks"]}, {"orders", "products"})

    def test_sees_changes_still_in_the_wal(self):
        writer = sqlite3.connect(self.path)
        writer.execute("PRAGMA journal_mode=WAL")
        writer.execute("PRAGMA wal_autocheckpoint=0")
        writer.execute("INSERT INTO regions VALUES (99, 'Only in WAL')")
        writer.commit()  # left open so the change isn't checkpointed
        conn = dbexplorer.connect(self.path)
        self.assertEqual(conn.execute("SELECT name FROM regions WHERE id = 99").fetchone()[0], "Only in WAL")
        writer.close()

    def test_paths_are_cleaned(self):
        home = os.path.expanduser("~")
        self.assertEqual(dbexplorer.clean_path("'~/a b.db'"), os.path.join(home, "a b.db"))
        if os.name == "nt":
            # File Explorer's "Copy as path", including network shares
            self.assertEqual(dbexplorer.clean_path('"C:\\My Data\\a.db"'), "C:\\My Data\\a.db")
            self.assertEqual(dbexplorer.clean_path('"\\\\server\\share\\a.db"'), "\\\\server\\share\\a.db")
            self.assertEqual(dbexplorer.clean_path("file:///C:/data/a%20b.db"), "C:\\data\\a b.db")
        else:
            # dragged into a Mac Terminal
            self.assertEqual(dbexplorer.clean_path("~/a\\ b.db"), os.path.join(home, "a b.db"))
            self.assertEqual(dbexplorer.clean_path("file:///tmp/a%20b.db"), "/tmp/a b.db")

    def test_not_a_database(self):
        with self.assertRaisesRegex(ValueError, "not a SQLite database"):
            dbexplorer.check_database(os.path.join(ROOT, "index.html"))


class Server(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = dbexplorer.start_server(18765)
        cls.base = f"http://127.0.0.1:{cls.server.server_address[1]}/"
        threading.Thread(target=cls.server.serve_forever, daemon=True).start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()

    def call(self, path, body=None, headers=None):
        req = urllib.request.Request(self.base + path, method="POST" if body is not None else "GET",
                                     data=None if body is None else json.dumps(body).encode(),
                                     headers={"X-DBX": "1", "Content-Type": "application/json", **(headers or {})})
        with urllib.request.urlopen(req) as r:
            data = r.read()
            return json.loads(data) if r.headers.get("Content-Type") == "application/json" else data

    def test_api_needs_header_and_host(self):
        for headers in [{"X-DBX": ""}, {"Host": "evil.example"}]:
            with self.assertRaises(urllib.error.HTTPError) as e:
                self.call("api/startup", headers=headers)
            self.assertEqual(e.exception.code, 403)

    def test_page_scripts_are_never_served_stale(self):
        req = urllib.request.Request(self.base)
        with urllib.request.urlopen(req) as r:
            html = r.read().decode()
            self.assertEqual(r.headers.get("Cache-Control"), "no-cache")
        for f in ["js/dom.js", "js/pipeline.js", "js/app.js", "css/app.css"]:
            self.assertRegex(html, re.escape(f) + r"\?v=\d+", f)
        with urllib.request.urlopen(self.base + "js/dom.js?v=1") as r:
            self.assertEqual(r.headers.get("Cache-Control"), "no-cache")
            self.assertIn(b"function icon", r.read())

    def test_only_app_files_are_served(self):
        for path in ["dbexplorer.py", ".git/config", "examples/sample.sqlite", "js/../dbexplorer.py"]:
            with self.assertRaises(urllib.error.HTTPError, msg=path):
                self.call(path)

    def test_open_query_export_download(self):
        db = self.call("api/open", {"path": SAMPLE})
        r = self.call("api/query", {"id": db["id"], "key": "page", "limit": 2,
                                    "sql": "SELECT id, loyalty_card, postcode FROM customers ORDER BY id"})
        self.assertEqual(len(r["rows"]), 2)
        self.assertTrue(r["more"])
        self.assertEqual(r["rows"][0][1], {"$int": "9007199254741007"})  # too big for JS numbers

        job = self.call("api/export", {"id": db["id"], "filename": "c/d", "sheets": [
            {"name": "customers", "sql": "SELECT * FROM customers WHERE id <= ?", "params": [10]}]})
        for _ in range(100):
            s = self.call(f"api/export/{job['token']}")
            if s["state"] != "running":
                break
            time.sleep(0.05)
        self.assertEqual(s["state"], "done", s)
        self.assertEqual(s["rows"], 10)
        self.assertEqual(s["filename"], "c_d.xlsx")
        data = self.call(f"download/{job['token']}", headers={"X-DBX": ""})
        if openpyxl:
            ws = openpyxl.load_workbook(io.BytesIO(data)).active
            self.assertEqual(ws.max_row, 11)
        with self.assertRaises(urllib.error.HTTPError):  # downloads are one-time
            self.call(f"download/{job['token']}")

    def test_new_query_interrupts_the_old_one(self):
        db = self.call("api/open", {"path": SAMPLE})
        slow = ("WITH RECURSIVE c(x) AS (SELECT 1 UNION ALL SELECT x + 1 FROM c) "
                "SELECT COUNT(*) FROM c WHERE x < 1e12")  # would run for a very long time
        result = {}

        def run():
            try:
                self.call("api/query", {"id": db["id"], "key": "count", "sql": slow})
            except urllib.error.HTTPError as e:
                result["code"] = e.code

        t = threading.Thread(target=run)
        t.start()
        time.sleep(0.3)
        r = self.call("api/query", {"id": db["id"], "key": "count", "sql": "SELECT 1"})
        t.join(5)
        self.assertFalse(t.is_alive())
        self.assertEqual(result.get("code"), 409)
        self.assertEqual(r["rows"], [[1]])


if __name__ == "__main__":
    unittest.main()
